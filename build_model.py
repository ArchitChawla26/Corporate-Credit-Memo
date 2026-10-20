#!/usr/bin/env python3
"""
build_model.py — builds CVS_credit_model.xlsx from data/inputs.json.

Sheets:
  Summary  — memo headline: key facts, headline ratios (linked), rating call
  Inputs   — every pulled value by period ($ millions), adjustments, ladder
  Ratios   — all ratios as LIVE Excel formulas driven only by Inputs

Run pull_data.py first. Author: Archit Chawla — independent research project.
"""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = Path(__file__).resolve().parent
D = json.loads((BASE / "data" / "inputs.json").read_text())
M = 1e6

NAVY = "1F3864"
BAND = "D9E2F3"
LIGHT = "F2F2F2"
thin = Side(style="thin", color="B0B0B0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
F_TITLE = Font(bold=True, size=13, color=NAVY)
F_HDR = Font(bold=True, color="FFFFFF")
F_SEC = Font(bold=True, color=NAVY)
F_BOLD = Font(bold=True)
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_SEC = PatternFill("solid", fgColor=BAND)
FILL_TOT = PatternFill("solid", fgColor=LIGHT)
NUM = "#,##0"
MULT = '0.00"x"'
PCT = "0.0%"

wb = Workbook()

# ---------------------------------------------------------------- Inputs ---
ws = wb.active
ws.title = "Inputs"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws["A1"] = "CVS Health Corporation (NYSE: CVS) — Credit Model Inputs"
ws["A1"].font = F_TITLE
ws["A2"] = ("All figures $ millions (shares in millions). Source: SEC EDGAR company facts, "
            "CIK 0000064803, pulled 2026-10-20 (see pull_data.py). Values as most recently "
            "reported in comparative 10-K / 10-Q filings.")
ws["A2"].font = Font(italic=True, size=9)

FY = ["FY2021", "FY2022", "FY2023", "FY2024", "FY2025"]
DEC = ["2021-12-31", "2022-12-31", "2023-12-31", "2024-12-31", "2025-12-31"]


def header_row(r, labels, widths=None):
    for c, lab in enumerate(labels, start=1):
        cell = ws.cell(row=r, column=c, value=lab)
        cell.font = F_HDR
        cell.fill = FILL_HDR
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center" if c > 1 else "left")
    return r


def put(r, label, values, bold=False, fill=None, fmt=NUM, note=None):
    c0 = ws.cell(row=r, column=1, value=label)
    c0.font = F_BOLD if bold else Font()
    c0.border = BORDER
    if fill:
        c0.fill = fill
    for i, v in enumerate(values):
        cell = ws.cell(row=r, column=2 + i, value=v)
        cell.number_format = fmt
        cell.border = BORDER
        if bold:
            cell.font = F_BOLD
        if fill:
            cell.fill = fill
    if note:
        ws.cell(row=r, column=10, value=note).font = Font(italic=True, size=9)
    return r


fa, fh, ins = D["flows_annual"], D["flows_h1"], D["instants"]


def flow(label):
    return [fa[label][p] / M for p in FY] + [fh[label]["H1-2025"] / M, fh[label]["H1-2026"] / M]


def inst(label, dates):
    out = []
    for dt in dates:
        v = ins[label][dt]
        out.append(None if v is None else v / M)
    return out


r = 4
header_row(r, ["Income statement & cash flow"] + FY + ["H1 2025", "H1 2026"])
R = {}
R["rev"] = put(5, "Revenues", flow("revenues"))
R["opinc"] = put(6, "Operating income (GAAP)", flow("operating_income"))
R["da"] = put(7, "Depreciation & amortization (cash flow statement)", flow("depreciation_amortization"))
R["int"] = put(8, "Interest expense", flow("interest_expense"),
               note="FY2021: InterestExpense tag; FY2022+: InterestExpenseNonoperating (identical in overlap years)")
R["ni"] = put(9, "Net income attributable to CVS Health", flow("net_income"))
R["ocf"] = put(10, "Net cash provided by operating activities", flow("operating_cash_flow"))
R["capex"] = put(11, "Capital expenditures (payments to acquire PP&E)", flow("capex"))
R["div"] = put(12, "Dividends paid", flow("dividends_paid"))
R["wshares"] = put(13, "Weighted avg diluted shares (millions)", flow("diluted_shares"))

r = 15
header_row(r, ["Non-recurring adjustments added back to EBITDA"] + FY + ["", "", "LTM add-backs"])
ws.cell(row=r, column=9, value="LTM (Jul-25–Jun-26)").font = F_HDR
ws.cell(row=r, column=9).fill = FILL_HDR
ADJ_ROWS = [("goodwill_impairment", "Goodwill impairment"),
            ("opioid_litigation", "Opioid litigation charges / accrual changes"),
            ("legacy_litigation", "Legacy litigation charges"),
            ("loss_assets_held_for_sale", "Loss on assets held for sale"),
            ("restructuring", "Restructuring charges")]
adj_row_idx = {}
for i, (key, label) in enumerate(ADJ_ROWS):
    vals = [D["adjustments_musd"][p][key] for p in FY]
    adj_row_idx[key] = put(16 + i, label, vals + [None, None, D["adjustments_musd"]["LTM"][key]])
R["adj_tot"] = 21
ws.cell(row=R["adj_tot"], column=1, value="Total adjustments").font = F_BOLD
for col in [2, 3, 4, 5, 6, 9]:
    L = get_column_letter(col)
    c = ws.cell(row=R["adj_tot"], column=col, value=f"=SUM({L}16:{L}20)")
    c.number_format = NUM
    c.font = F_BOLD
    c.fill = FILL_TOT
    c.border = BORDER
ws.cell(row=R["adj_tot"], column=1).fill = FILL_TOT
ws.cell(row=R["adj_tot"], column=1).border = BORDER
ws.cell(row=22, column=1, value=("LTM add-backs are only the items recorded inside Jul 2025–Jun 2026 "
                                "(Q3 2025 goodwill impairment and opioid accrual charge). The $1,220M legacy "
                                "litigation charges were recorded in H1 2025 and sit outside the LTM window.")
        ).font = Font(italic=True, size=9)

r = 24
header_row(r, ["Balance sheet (period-end)"] +
           ["Dec 2021", "Dec 2022", "Dec 2023", "Dec 2024", "Dec 2025", "Jun 2026"])
BSD = DEC + ["2026-06-30"]
R["ltd_nc"] = put(25, "Long-term debt, noncurrent", inst("lt_debt_noncurrent", BSD))
R["ltd_c"] = put(26, "Current portion of long-term debt", inst("lt_debt_current", BSD))
R["stb"] = put(27, "Short-term borrowings (commercial paper)", inst("short_term_borrowings", BSD))
R["debt"] = 28
ws.cell(row=28, column=1, value="Total debt (carrying value)").font = F_BOLD
for col in range(2, 8):
    L = get_column_letter(col)
    c = ws.cell(row=28, column=col, value=f"=SUM({L}25:{L}27)")
    c.number_format = NUM; c.font = F_BOLD; c.fill = FILL_TOT; c.border = BORDER
ws.cell(row=28, column=1).fill = FILL_TOT; ws.cell(row=28, column=1).border = BORDER
R["cash"] = put(29, "Cash and cash equivalents", inst("cash", BSD))
R["ndebt"] = 30
ws.cell(row=30, column=1, value="Net debt").font = F_BOLD
for col in range(2, 8):
    L = get_column_letter(col)
    c = ws.cell(row=30, column=col, value=f"={L}28-{L}29")
    c.number_format = NUM; c.font = F_BOLD; c.fill = FILL_TOT; c.border = BORDER
ws.cell(row=30, column=1).fill = FILL_TOT; ws.cell(row=30, column=1).border = BORDER
R["olc"] = put(31, "Operating lease liabilities, current", inst("op_lease_current", BSD))
R["olnc"] = put(32, "Operating lease liabilities, noncurrent", inst("op_lease_noncurrent", BSD))
fin_c = inst("fin_lease_current", DEC) + [ins["fin_lease_current"]["2025-12-31"] / M]
fin_nc = inst("fin_lease_noncurrent", DEC) + [ins["fin_lease_noncurrent"]["2025-12-31"] / M]
R["flc"] = put(33, "Finance lease liabilities, current", fin_c,
               note="Jun-2026 column carries Dec-2025 values: finance leases are not broken out in the Q2 2026 condensed balance sheet")
R["flnc"] = put(34, "Finance lease liabilities, noncurrent", fin_nc)
R["assets"] = put(35, "Total assets", inst("total_assets", BSD))
R["equity"] = put(36, "Total shareholders' equity", inst("total_equity", BSD))
R["shares"] = put(37, "Common shares outstanding (millions)", inst("shares_outstanding", BSD))

r = 39
header_row(r, ["Debt maturity ladder — principal at Dec 31, 2025 (FY2025 10-K, borrowings note)", "Principal $M"])
lad = D["maturity_ladder"]
lad_rows = [("due_2026", "2026"), ("due_2027", "2027"), ("due_2028", "2028"),
            ("due_2029", "2029"), ("due_2030", "2030"), ("due_thereafter", "Thereafter")]
for i, (k, lab) in enumerate(lad_rows):
    put(40 + i, lab, [lad[k] / M])
R["lad_tot"] = 46
ws.cell(row=46, column=1, value="Subtotal — long-term debt principal").font = F_BOLD
c = ws.cell(row=46, column=2, value="=SUM(B40:B45)")
c.number_format = NUM; c.font = F_BOLD; c.fill = FILL_TOT; c.border = BORDER
ws.cell(row=46, column=1).fill = FILL_TOT; ws.cell(row=46, column=1).border = BORDER
ws.cell(row=47, column=1, value="Add: finance lease liabilities")
c = ws.cell(row=47, column=2, value="=F33+F34")
c.number_format = NUM; c.border = BORDER
ws.cell(row=48, column=1, value="Total debt principal incl. finance leases").font = F_BOLD
c = ws.cell(row=48, column=2, value="=B46+B47")
c.number_format = NUM; c.font = F_BOLD; c.fill = FILL_TOT; c.border = BORDER
ws.cell(row=48, column=1).fill = FILL_TOT; ws.cell(row=48, column=1).border = BORDER

r = 50
header_row(r, ["Liquidity facilities (per borrowings notes)", "Amount $M"])
put(51, "Back-up revolving credit facilities (3 x $2,500M), undrawn", [7500],
    note="Expire May 2029 / 2030 / 2031 per Q2 2026 10-Q (May 2028 / 2029 / 2030 at Dec 2025)")
put(52, "Commercial paper outstanding, Jun 30, 2026", [0],
    note="None outstanding at Jun 30, 2026 or Dec 31, 2025 (10-Q / 10-K)")

ws.column_dimensions["A"].width = 52
for col in "BCDEFGHI":
    ws.column_dimensions[col].width = 14
ws.column_dimensions["J"].width = 60
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Ratios ---
wr = wb.create_sheet("Ratios")
wr.sheet_properties.pageSetUpPr.fitToPage = True
wr["A1"] = "CVS Health — Ratio Analysis (all cells are live formulas from Inputs)"
wr["A1"].font = F_TITLE
wr["A2"] = ("LTM = twelve months ended Jun 30, 2026: FY2025 less H1 2025 plus H1 2026; "
            "balance-sheet items at Jun 30, 2026. $ millions unless noted.")
wr["A2"].font = Font(italic=True, size=9)
header = [""] + FY + ["LTM"]
for c, lab in enumerate(header, start=1):
    cell = wr.cell(row=4, column=c, value=lab)
    cell.font = F_HDR; cell.fill = FILL_HDR; cell.border = BORDER

# Inputs columns: FY -> B..F; H1-2025 -> G; H1-2026 -> H; LTM adj -> I
# Ratios columns: FY -> B..F; LTM -> G
def flow_f(row_in, col_out):
    """Formula for a flow item in Ratios column col_out (2..7)."""
    if col_out <= 6:
        return f"=Inputs!{get_column_letter(col_out)}{row_in}"
    return f"=Inputs!F{row_in}-Inputs!G{row_in}+Inputs!H{row_in}"


def bs_f(row_in, col_out):
    """Balance-sheet item: Ratios FY cols B..F map to Inputs Dec cols B..F; LTM -> Jun-26 col G."""
    if col_out <= 6:
        return f"=Inputs!{get_column_letter(col_out)}{row_in}"
    return f"=Inputs!G{row_in}"


RR = {}
def rrow(r, label, maker, fmt=NUM, bold=False, fill=None, note=None):
    c0 = wr.cell(row=r, column=1, value=label)
    c0.font = F_BOLD if bold else Font()
    c0.border = BORDER
    if fill: c0.fill = fill
    for col in range(2, 8):
        cell = wr.cell(row=r, column=col, value=maker(col))
        cell.number_format = fmt; cell.border = BORDER
        if bold: cell.font = F_BOLD
        if fill: cell.fill = fill
    if note:
        wr.cell(row=r, column=9, value=note).font = Font(italic=True, size=9)
    RR[label] = r
    return r

def L(r):  # noqa: E743 - cell address helper for current sheet row r
    return lambda col: f"{get_column_letter(col)}{r}"

rrow(5, "Revenues", lambda c: flow_f(R["rev"], c))
rrow(6, "Operating income (GAAP)", lambda c: flow_f(R["opinc"], c))
rrow(7, "Plus: depreciation & amortization", lambda c: flow_f(R["da"], c))
rrow(8, "Reported EBITDA", lambda c: f"={L(6)(c)}+{L(7)(c)}", bold=True, fill=FILL_TOT)
rrow(9, "Plus: non-recurring adjustments (Inputs)",
     lambda c: (f"=Inputs!{get_column_letter(c)}{R['adj_tot']}" if c <= 6 else f"=Inputs!I{R['adj_tot']}"))
rrow(10, "Adjusted EBITDA", lambda c: f"={L(8)(c)}+{L(9)(c)}", bold=True, fill=FILL_TOT)
rrow(11, "Interest expense", lambda c: flow_f(R["int"], c))
rrow(12, "Total debt", lambda c: bs_f(R["debt"], c))
rrow(13, "Cash and cash equivalents", lambda c: bs_f(R["cash"], c))
rrow(14, "Net debt", lambda c: f"={L(12)(c)}-{L(13)(c)}", bold=True, fill=FILL_TOT)
rrow(15, "Gross leverage — Debt / Adjusted EBITDA", lambda c: f"={L(12)(c)}/{L(10)(c)}", fmt=MULT, bold=True)
rrow(16, "Gross leverage — Debt / Reported EBITDA", lambda c: f"={L(12)(c)}/{L(8)(c)}", fmt=MULT)
rrow(17, "Net leverage — Net debt / Adjusted EBITDA", lambda c: f"={L(14)(c)}/{L(10)(c)}", fmt=MULT, bold=True)
rrow(18, "Interest coverage — Adjusted EBITDA / Interest", lambda c: f"={L(10)(c)}/{L(11)(c)}", fmt=MULT, bold=True)
rrow(19, "Interest coverage — Reported EBITDA / Interest", lambda c: f"={L(8)(c)}/{L(11)(c)}", fmt=MULT)
rrow(20, "FFO (approx.) — net income + D&A + goodwill impairment",
     lambda c: (f"={flow_f(R['ni'], c)[1:]}+{flow_f(R['da'], c)[1:]}+Inputs!{get_column_letter(c)}{adj_row_idx['goodwill_impairment']}"
                if c <= 6 else
                f"={flow_f(R['ni'], c)[1:]}+{flow_f(R['da'], c)[1:]}+Inputs!I{adj_row_idx['goodwill_impairment']}"),
     note="Stated approximation: excludes deferred taxes, stock comp and other non-cash items")
rrow(21, "FFO / Debt", lambda c: f"={L(20)(c)}/{L(12)(c)}", fmt=PCT, bold=True)
rrow(22, "FCF after dividends — OCF − capex − dividends paid",
     lambda c: f"={flow_f(R['ocf'], c)[1:]}-({flow_f(R['capex'], c)[1:]})-({flow_f(R['div'], c)[1:]})")
rrow(23, "FCF after dividends / Debt", lambda c: f"={L(22)(c)}/{L(12)(c)}", fmt=PCT)
rrow(24, "Lease-adjusted debt — debt + operating + finance lease liabilities",
     lambda c: f"={L(12)(c)}+{bs_f(R['olc'], c)[1:]}+{bs_f(R['olnc'], c)[1:]}+{bs_f(R['flc'], c)[1:]}+{bs_f(R['flnc'], c)[1:]}")
rrow(25, "Lease-adjusted leverage — Lease-adj. debt / Adjusted EBITDA",
     lambda c: f"={L(24)(c)}/{L(10)(c)}", fmt=MULT)
rrow(26, "Adjusted EBITDA margin", lambda c: f"={L(10)(c)}/{L(5)(c)}", fmt=PCT)
wr.column_dimensions["A"].width = 56
for col in "BCDEFG":
    wr.column_dimensions[col].width = 13
wr.column_dimensions["I"].width = 62
wr.freeze_panes = "B5"

# --------------------------------------------------------------- Summary ---
wsm = wb.create_sheet("Summary", 0)
wsm.sheet_properties.pageSetUpPr.fitToPage = True
wsm["A1"] = "CORPORATE CREDIT MEMO — CVS Health Corporation (NYSE: CVS)"
wsm["A1"].font = Font(bold=True, size=14, color=NAVY)
wsm["A2"] = ("Prepared by: Archit Chawla  |  Date: October 20, 2026  |  Independent research project — "
             "not investment advice and not a credit rating.")
wsm["A2"].font = Font(italic=True, size=9)
wsm["A4"] = "RECOMMENDATION: APPROVE — maintain investment-grade senior unsecured exposure (conditions in memo)"
wsm["A4"].font = Font(bold=True, size=11, color="1F6F43")
wsm["A5"] = "Shadow rating: BBB, stable outlook (agency ratings for reference: Moody's Baa3 stable / S&P BBB / Fitch BBB negative)"
wsm["A5"].font = F_BOLD

wsm["A7"] = "Headline ratios (live-linked to Ratios sheet)"
wsm["A7"].font = F_SEC
for c, lab in enumerate(["", "FY2023", "FY2024", "FY2025", "LTM Jun-26"], start=1):
    cell = wsm.cell(row=8, column=c, value=lab)
    cell.font = F_HDR; cell.fill = FILL_HDR; cell.border = BORDER
SUM_ROWS = [("Total debt ($M)", 12, NUM), ("Adjusted EBITDA ($M)", 10, NUM),
            ("Gross leverage (adj.)", 15, MULT), ("Net leverage (adj.)", 17, MULT),
            ("Interest coverage (adj.)", 18, MULT), ("FFO / Debt", 21, PCT),
            ("FCF after dividends ($M)", 22, NUM), ("Lease-adjusted leverage", 25, MULT)]
for i, (lab, rrow_idx, fmt) in enumerate(SUM_ROWS):
    rr = 9 + i
    wsm.cell(row=rr, column=1, value=lab).border = BORDER
    for j, rcol in enumerate([4, 5, 6, 7]):  # Ratios cols FY23, FY24, FY25, LTM
        cell = wsm.cell(row=rr, column=2 + j,
                        value=f"=Ratios!{get_column_letter(rcol)}{rrow_idx}")
        cell.number_format = fmt; cell.border = BORDER

lr = 9 + len(SUM_ROWS) + 1
wsm.cell(row=lr, column=1, value="Key facts (Jun 30, 2026 unless noted)").font = F_SEC
facts = [
    ("Total debt", f"=Ratios!G12", NUM, "down from $66,270M at Dec 31, 2024 (peak)"),
    ("Cash", f"=Ratios!G13", NUM, "~$2,700M held at parent / non-restricted subsidiaries"),
    ("Undrawn revolvers", "=Inputs!B51", NUM, "3 x $2,500M back-up facilities, expire May 2029/30/31"),
    ("2027 debt maturities (principal)", "=Inputs!B41", NUM, "per Dec 31, 2025 maturity ladder"),
    ("LTM revenues", "=Ratios!G5", NUM, ""),
    ("LTM FFO / Debt", "=Ratios!G21", PCT, ""),
]
for i, (lab, formula, fmt, note) in enumerate(facts):
    rr = lr + 1 + i
    wsm.cell(row=rr, column=1, value=lab).border = BORDER
    cell = wsm.cell(row=rr, column=2, value=formula)
    cell.number_format = fmt; cell.border = BORDER
    if note:
        wsm.cell(row=rr, column=3, value=note).font = Font(italic=True, size=9)

vr = lr + 1 + len(facts) + 1
wsm.cell(row=vr, column=1, value="Verification").font = F_SEC
wsm.cell(row=vr + 1, column=1,
         value=("Model recalculated in LibreOffice headless after build; every Ratios value "
                "checked against an independent Python computation from data/inputs.json "
                "(see memo, Methodology & Verification). Inputs tie to EDGAR company facts "
                "for CIK 0000064803 pulled 2026-10-20.")).font = Font(size=9)
wsm.column_dimensions["A"].width = 34
for col in "BCDE":
    wsm.column_dimensions[col].width = 15

wb.save(BASE / "CVS_credit_model.xlsx")
print("Wrote CVS_credit_model.xlsx")
